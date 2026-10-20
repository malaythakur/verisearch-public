"""Text extraction from uploaded files (PDF, DOCX, TXT/MD, HTML).

Used by the upload endpoint to turn a user's document into the cleaned text that the
indexer ingests. Each extractor is best-effort and raises ExtractionError on a format
it cannot handle so the API can return a clear 422.
"""

from __future__ import annotations

import io

# Cap accepted upload size (10 MB) to bound memory and abuse.
MAX_UPLOAD_BYTES = 10 * 1024 * 1024

# Supported content by file extension.
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown", ".html", ".htm"}


class ExtractionError(Exception):
    """Raised when a file's text cannot be extracted."""

    def __init__(self, message: str, *, code: str = "extraction_failed") -> None:
        self.code = code
        self.message = message
        super().__init__(message)


def _ext(filename: str) -> str:
    name = (filename or "").lower()
    dot = name.rfind(".")
    return name[dot:] if dot != -1 else ""


def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from an uploaded file's bytes.

    Args:
        filename: Original filename (used to detect format by extension).
        data: Raw file bytes.

    Returns:
        Extracted plain text.

    Raises:
        ExtractionError: on unsupported format, oversize, or parse failure.
    """
    if len(data) > MAX_UPLOAD_BYTES:
        raise ExtractionError(
            f"file exceeds {MAX_UPLOAD_BYTES // (1024 * 1024)}MB limit",
            code="file_too_large",
        )

    ext = _ext(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(
            f"unsupported file type '{ext or 'unknown'}'. Supported: "
            f"{', '.join(sorted(SUPPORTED_EXTENSIONS))}",
            code="unsupported_file_type",
        )

    try:
        if ext == ".pdf":
            return _extract_pdf(data)
        if ext == ".docx":
            return _extract_docx(data)
        # txt, md, html — decode as text (HTML is cleaned downstream by the indexer)
        return data.decode("utf-8", errors="replace")
    except ExtractionError:
        raise
    except Exception as e:  # noqa: BLE001
        raise ExtractionError(f"failed to extract text: {e}") from e


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text)
    result = "\n\n".join(parts).strip()
    if not result:
        raise ExtractionError(
            "no extractable text in PDF (it may be scanned/image-only)",
            code="empty_extraction",
        )
    return result


def _extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    # Include table cell text too.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text and cell.text.strip():
                    parts.append(cell.text)
    result = "\n".join(parts).strip()
    if not result:
        raise ExtractionError("no extractable text in DOCX", code="empty_extraction")
    return result
